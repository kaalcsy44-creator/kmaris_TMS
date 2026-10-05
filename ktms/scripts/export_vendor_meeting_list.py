"""공급사 미팅 후보 목록(xlsx) — 견적이 실제로 오간 공급사만 추려 등록 정보·취급 품목·
메이커·견적 횟수·견적 품목·발주/거래성사 여부를 한 파일로 낸다.

"견적이 오갔다" = 그 회사 앞으로 받은 벤더 견적(vendor_quotes)이 1건 이상. RFQ 만 보내고
회신이 없던 곳은 본 목록에서 빼고 별도 시트(참고_RFQ만발송)에 둔다.

거래선은 레코드 1건 = 담당자 1명이라 같은 회사가 여러 줄이다 — 회사 이름으로 묶는다.
부산 소재(주소·지역에 부산/Busan) 업체를 맨 위로 올린다.

사용:
    DATABASE_URL=postgresql://... python ktms/scripts/export_vendor_meeting_list.py [출력.xlsx]
"""
from __future__ import annotations

import json
import os
import re
import sys
from collections import defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import create_engine, text

FONT = "Noto Sans KR"
BUSAN = ("부산", "busan", "pusan")
NEAR = ("김해", "gimhae", "양산", "yangsan", "창원", "changwon", "진해", "jinhae", "거제", "geoje")
RANK = {"Y": 0, "인근": 1, "": 2}


def _j(v, default):
    """JSON 칸 — Postgres 는 파싱된 값, SQLite 는 문자열로 온다."""
    if v is None or v == "":
        return default
    if isinstance(v, (list, dict)):
        return v
    try:
        return json.loads(v)
    except Exception:
        return default


def _norm(name: str) -> str:
    return re.sub(r"[\s.,()\-]+", "", (name or "").lower())


def _f(v):
    try:
        return float(str(v).replace(",", ""))
    except Exception:
        return None


def _join(vals, sep=" / "):
    out = []
    for v in vals:
        v = str(v or "").strip()
        if v and v not in out:
            out.append(v)
    return sep.join(out)


def _is_item(it: dict) -> bool:
    if not isinstance(it, dict):
        return False
    if str(it.get("row_kind") or "") == "option":
        return False
    return bool(str(it.get("part_no") or "").strip() or str(it.get("description") or "").strip())


def load(url: str):
    if url.startswith("postgres://"):
        url = url.replace("postgres://", "postgresql://", 1)
    eng = create_engine(url)
    q = lambda c, s: [dict(r._mapping) for r in c.execute(text(s))]  # noqa: E731
    with eng.connect() as c:
        # select * — 열이 덜 붙은 옛 DB(로컬 SQLite)에서도 돌도록. 없는 표는 빈 목록.
        t = lambda name: q(c, f"select * from {name}") if eng.dialect.has_table(c, name) else []  # noqa: E731
        return {"vendors": t("vendors"), "makers": t("makers"), "cats": t("item_categories"),
                "customers": t("customers"), "vessels": t("vessels"), "rfqs": t("rfqs"),
                "vrfqs": t("vendor_rfqs"), "vqs": t("vendor_quotes"), "qtns": t("quotations"),
                "orders": t("orders"), "pos": t("purchase_orders"),
                "awards": t("deal_line_awards")}


def build(d: dict):
    vendors = {v["id"]: v for v in d["vendors"]}
    makers = {m["id"]: m["name"] for m in d["makers"]}
    cats = {c["id"]: c for c in d["cats"]}
    customers = {c["id"]: c["name"] for c in d["customers"]}
    vessels = {v["id"]: v["name"] for v in d["vessels"]}
    rfqs = {r["id"]: r for r in d["rfqs"]}
    vrfqs = {r["id"]: r for r in d["vrfqs"]}
    qtns = {q["id"]: q for q in d["qtns"]}
    orders = {o["id"]: o for o in d["orders"]}

    def cat_label(cid):
        c = cats.get(cid)
        if not c:
            return ""
        return f"{c['code']} {c['name']}" if c.get("code") else c.get("name")

    # 회사 단위 묶음(이름 정규화)
    comp_of = {vid: _norm(v["name"]) for vid, v in vendors.items()}

    # 딜(RFQ) → 고객 수주 여부
    ordered_rfqs = set()
    for o in orders.values():
        rid = o.get("rfq_id") or (qtns.get(o.get("quotation_id")) or {}).get("rfq_id")
        if rid:
            ordered_rfqs.add(rid)
    for r in rfqs.values():
        if str(r.get("status") or "").upper() in ("ORDERED", "수주완료"):
            ordered_rfqs.add(r["id"])

    # 채택 — 라인 채택 + 견적서가 원가 출처로 고른 벤더 견적 + 줄의 src_vq_id
    awarded_vq = defaultdict(set)   # vq_id -> {lid or "*"}
    for a in d["awards"]:
        if a.get("vendor_quote_id"):
            awarded_vq[a["vendor_quote_id"]].add(a.get("lid") or "*")
    for q in qtns.values():
        if q.get("vendor_quote_id"):
            awarded_vq[q["vendor_quote_id"]].add("*")
        for it in _j(q.get("items"), []):
            if isinstance(it, dict) and it.get("src_vq_id"):
                try:
                    awarded_vq[int(it["src_vq_id"])].add(it.get("lid") or "*")
                except Exception:
                    pass

    # P/O — 회사·딜 단위
    po_by_comp = defaultdict(list)
    po_deals = defaultdict(set)     # comp -> rfq ids
    po_rows = []
    for p in d["pos"]:
        v = vendors.get(p.get("vendor_id"))
        if not v:
            continue
        comp = comp_of[v["id"]]
        o = orders.get(p.get("order_id")) or {}
        rid = o.get("rfq_id") or (qtns.get(o.get("quotation_id")) or {}).get("rfq_id")
        items = [it for it in _j(p.get("items"), []) if _is_item(it)]
        amt = 0.0
        for it in items:
            a = _f(it.get("amount"))
            if a is None:
                a = (_f(it.get("qty")) or 0) * (_f(it.get("unit_price")) or 0)
            amt += a or 0
        po_by_comp[comp].append((p, amt))
        if rid:
            po_deals[comp].add(rid)
        r = rfqs.get(rid) or {}
        po_rows.append({
            "업체명": v["name"], "P/O No.": p.get("po_no"), "발주일": p.get("date"),
            "딜(RFQ No.)": r.get("rfq_no"), "프로젝트": r.get("project_title"),
            "고객": customers.get(r.get("customer_id")), "고객 P/O": o.get("po_no"),
            "품목": _join([f"{it.get('part_no') or ''} {it.get('description') or ''}".strip() for it in items], "\n"),
            "품목 수": len(items), "통화": p.get("currency"), "발주금액": round(amt, 2) or None,
            "상태": p.get("status"),
        })

    # 견적 이력
    comp_stat = defaultdict(lambda: {"vq": 0, "deals": set(), "parts": [], "makers": [], "dates": [],
                                     "awarded": 0, "ordered_deals": set(), "vrfq": 0})
    for vr in vrfqs.values():
        if vr.get("vendor_id") in vendors:
            comp_stat[comp_of[vr["vendor_id"]]]["vrfq"] += 1
    detail_rows = []
    for vq in d["vqs"]:
        vr = vrfqs.get(vq.get("vendor_rfq_id"))
        if not vr or vr.get("vendor_id") not in vendors:
            continue
        v = vendors[vr["vendor_id"]]
        comp = comp_of[v["id"]]
        st = comp_stat[comp]
        st["vq"] += 1
        rid = vr.get("rfq_id")
        r = rfqs.get(rid) or {}
        if rid:
            st["deals"].add(rid)
            if rid in ordered_rfqs:
                st["ordered_deals"].add(rid)
        if vq.get("received_date"):
            st["dates"].append(vq["received_date"])
        aw = awarded_vq.get(vq["id"], set())
        items = [it for it in _j(vq.get("items"), []) if _is_item(it)]
        terms = _j(vq.get("terms"), {})
        for it in items:
            pn = str(it.get("part_no") or "").strip()
            desc = str(it.get("description") or "").strip()
            mk = str(it.get("maker") or it.get("manufacturer") or "").strip()
            st["parts"].append(desc or pn)
            if mk:
                st["makers"].append(mk)
            is_aw = "*" in aw or (it.get("lid") and it.get("lid") in aw)
            if is_aw:
                st["awarded"] += 1
            up = _f(it.get("cost_price") if it.get("cost_price") not in (None, "") else it.get("unit_price"))
            qty = _f(it.get("qty"))
            amt = _f(it.get("amount"))
            if amt is None and up is not None and qty is not None:
                amt = up * qty
            detail_rows.append({
                "업체명": v["name"], "딜(RFQ No.)": r.get("rfq_no"), "프로젝트": r.get("project_title"),
                "고객": customers.get(r.get("customer_id")), "선박": vessels.get(r.get("vessel_id")),
                "벤더 RFQ No.": vr.get("kmaris_rfq_no"), "RFQ 발송일": vr.get("sent_date"),
                "벤더 견적번호": vq.get("vendor_quote_no"), "견적 수신일": vq.get("received_date"),
                "품번": pn, "품명": desc, "메이커": mk, "수량": qty, "단위": it.get("unit"),
                "통화": vq.get("currency"), "단가": up, "금액": amt,
                "납기": it.get("lead_time") or terms.get("lead_time") or terms.get("delivery"),
                "고객견적 반영": "Y" if is_aw else "",
                "고객 수주(딜)": "Y" if rid in ordered_rfqs else "",
                "당사 발주(딜)": "Y" if rid in po_deals[comp] else "",
                "딜 상태": r.get("status"),
            })

    # 회사 요약
    groups = defaultdict(list)
    for vid, v in vendors.items():
        groups[comp_of[vid]].append(v)

    def busan_flag(recs):
        """주소가 우선이다 — 지역 칸은 '부산권'으로 넓게 적힌 경우가 있다(김해 주소에 Busan 지역)."""
        addr = " ".join(" ".join([str(r.get("address") or ""), " ".join(map(str, _j(r.get("addresses"), [])))])
                        for r in recs).lower()
        if not addr.strip():
            addr = " ".join(" ".join([str(r.get("country") or ""), " ".join(map(str, _j(r.get("regions"), [])))])
                            for r in recs).lower()
        if any(k in addr for k in BUSAN):
            return "Y"
        if any(k in addr for k in NEAR):
            return "인근"
        return ""

    summary, rfq_only, contacts = [], [], []
    for comp, recs in groups.items():
        recs.sort(key=lambda r: r["id"])
        st = comp_stat.get(comp)
        head = recs[0]
        cat_ids = {c for r in recs for c in _j(r.get("category_ids"), [])}
        mk_ids = {m for r in recs for m in _j(r.get("maker_ids"), [])}
        base = {
            "부산 소재": busan_flag(recs),
            "업체명": head["name"],
            "국가/지역": _join([x for r in recs for x in (_j(r.get("regions"), []) or [r.get("country")])]),
            "주소": _join([x for r in recs for x in (_j(r.get("addresses"), []) or [r.get("address")])], "\n"),
            "담당자": _join([f"{r.get('contact') or ''}" + (f" ({r['duty']})" if r.get("duty") else "") for r in recs], "\n"),
            "이메일": _join([x for r in recs for x in (_j(r.get("emails"), []) or [r.get("email")])], "\n"),
            "전화": _join([x for r in recs for x in (_j(r.get("phones"), []) or [r.get("contact_phone")])], "\n"),
            "웹사이트": _join([r.get("website") for r in recs]),
            "전문분야(등록)": _join([r.get("specialization") for r in recs], "\n"),
            "회사소개": _join([r.get("note") for r in recs], "\n"),
            "결제조건": _join([r.get("payment_terms") for r in recs]),
            "취급 품목분류(등록)": _join(sorted(cat_label(c) for c in cat_ids if cat_label(c)), "\n"),
            "공급 메이커(등록)": _join(sorted(makers.get(m, "") for m in mk_ids), ", "),
            "메이커 직거래": _join([makers.get(r.get("maker_id")) for r in recs if r.get("maker_id")]),
        }
        for r in recs:
            contacts.append({"업체명": r["name"], "담당자": r.get("contact"), "담당분야": r.get("duty"),
                             "이메일": _join(_j(r.get("emails"), []) or [r.get("email")], "\n"),
                             "전화": _join(_j(r.get("phones"), []) or [r.get("contact_phone")], "\n"),
                             "부산 소재": base["부산 소재"],
                             "견적 회신": "Y" if st and st["vq"] else ""})
        if not st or st["vq"] == 0:
            if st and st["vrfq"]:
                rfq_only.append({**base, "RFQ 수신 횟수": st["vrfq"]})
            continue
        pos = po_by_comp.get(comp, [])
        po_amt = defaultdict(float)
        for p, a in pos:
            po_amt[p.get("currency") or "?"] += a
        summary.append({
            "No": None, **base,
            "RFQ 수신 횟수": st["vrfq"],
            "견적 회신 횟수": st["vq"],
            "견적 딜 수": len(st["deals"]),
            "견적 품목 수": len(st["parts"]),
            "견적 품목": _join(st["parts"], "\n"),
            "견적 메이커": _join(st["makers"], ", "),
            "최근 견적일": max(st["dates"]) if st["dates"] else None,
            "고객견적 반영 품목 수": st["awarded"],
            "고객 수주된 딜 수": len(st["ordered_deals"]),
            "발주(P/O) 건수": len(pos),
            "발주금액": _join([f"{c} {a:,.2f}" for c, a in po_amt.items()]),
            "발주 여부": "Y" if pos else "N",
            "거래성사 여부": "성사" if pos else ("고객견적 반영(미발주)" if st["awarded"] else "미성사"),
        })

    summary.sort(key=lambda r: (RANK[r["부산 소재"]], -r["견적 회신 횟수"], r["업체명"]))
    for i, r in enumerate(summary, 1):
        r["No"] = i
    rfq_only.sort(key=lambda r: (RANK[r["부산 소재"]], r["업체명"]))
    order = {r["업체명"]: i for i, r in enumerate(summary)}
    detail_rows.sort(key=lambda r: (order.get(r["업체명"], 9999), str(r["견적 수신일"] or "")))
    po_rows.sort(key=lambda r: (order.get(r["업체명"], 9999), str(r["발주일"] or "")))
    contacts = [c for c in contacts if c["견적 회신"]]
    contacts.sort(key=lambda r: (order.get(r["업체명"], 9999)))
    return summary, detail_rows, po_rows, contacts, rfq_only


WIDTH = {"No": 5, "부산 소재": 7, "업체명": 26, "국가/지역": 12, "주소": 40, "담당자": 20, "이메일": 28, "전화": 18,
         "웹사이트": 24, "전문분야(등록)": 36, "회사소개": 40, "결제조건": 16, "취급 품목분류(등록)": 30,
         "공급 메이커(등록)": 30, "메이커 직거래": 14, "견적 품목": 40, "견적 메이커": 26, "품목": 40,
         "품명": 34, "품번": 18, "프로젝트": 24, "고객": 22, "선박": 18, "딜 상태": 12, "발주금액": 18}


def write_sheet(wb, title, rows, freeze="C2", note=None):
    ws = wb.create_sheet(title)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    if not rows:
        ws["A1"] = "해당 없음"
        ws["A1"].font = Font(name=FONT)
        return ws
    cols = list(rows[0].keys())
    for j, c in enumerate(cols, 1):
        cell = ws.cell(1, j, c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        ws.column_dimensions[get_column_letter(j)].width = WIDTH.get(c, 12)
    busan_fill = PatternFill("solid", fgColor="FFF2CC")
    near_fill = PatternFill("solid", fgColor="EDEDED")
    for i, r in enumerate(rows, 2):
        hl = {"Y": busan_fill, "인근": near_fill}.get(r.get("부산 소재"))
        for j, c in enumerate(cols, 1):
            v = r[c]
            cell = ws.cell(i, j, v if v != "" else None)
            cell.font = Font(name=FONT, size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
            if isinstance(v, float):
                cell.number_format = "#,##0.00"
            if hl:
                cell.fill = hl
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"
    return ws


def main():
    url = os.environ.get("DATABASE_URL")
    if not url:
        sys.exit("DATABASE_URL 이 필요합니다.")
    out = sys.argv[1] if len(sys.argv) > 1 else f"KTMS_공급사_미팅후보_{date.today():%Y-%m-%d}.xlsx"
    summary, detail, pos, contacts, rfq_only = build(load(url))

    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "미팅후보_요약", summary)
    write_sheet(wb, "견적이력_상세", detail)
    write_sheet(wb, "발주이력", pos)
    write_sheet(wb, "담당자", contacts)
    write_sheet(wb, "참고_RFQ만발송", rfq_only)
    ws = wb.create_sheet("기준")
    legend = [
        ("기준일", f"{date.today():%Y-%m-%d} (KTMS 운영 DB 추출)"),
        ("대상", "벤더 견적(Vendor Quote)을 1건 이상 받은 공급사. RFQ만 보내고 회신 없는 곳은 '참고_RFQ만발송' 시트"),
        ("업체 묶음", "거래선은 담당자마다 한 줄이라 같은 회사명(공백·기호 무시)을 한 회사로 합침"),
        ("부산 소재", "Y = 주소가 부산(노란 바탕) · 인근 = 김해·양산·창원·진해·거제(회색 바탕). 주소 기준, 주소가 없으면 등록 지역 기준. 요약 시트 상단 정렬"),
        ("RFQ 수신 횟수", "당사가 그 업체에 보낸 벤더 RFQ 건수"),
        ("견적 회신 횟수", "그 업체에서 받은 견적서 건수(재견적 포함)"),
        ("고객견적 반영", "3단계 라인 채택, 또는 고객 견적서의 원가 출처로 쓰인 벤더 견적/줄"),
        ("단가·금액", "공급사가 제시한 원가(cost price), 견적 통화 기준"),
        ("고객 수주된 딜", "그 업체가 견적한 딜 중 고객 P/O(오더)가 등록된 딜 — 다른 업체에서 샀을 수도 있음"),
        ("발주 여부", "당사가 그 업체에 P/O를 발행했는가"),
        ("거래성사 여부", "성사 = 당사 발주(P/O) 있음 · 고객견적 반영(미발주) = 고객 견적에 쓰였으나 P/O 미발행 · 미성사 = 견적만 받음"),
        ("등록 항목", "주소·담당자·전문분야·회사소개·취급 품목분류·공급 메이커는 KTMS 거래선 명부에 등록된 값"),
    ]
    for i, (k, v) in enumerate(legend, 1):
        ws.cell(i, 1, k).font = Font(name=FONT, bold=True)
        ws.cell(i, 2, v).font = Font(name=FONT)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 100
    wb.save(out)
    print(f"saved {out}: 견적 공급사 {len(summary)}곳(부산 {sum(r['부산 소재'] == 'Y' for r in summary)}), "
          f"견적 줄 {len(detail)}, P/O {len(pos)}, RFQ만 {len(rfq_only)}")


if __name__ == "__main__":
    main()
